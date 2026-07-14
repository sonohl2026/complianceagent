import io

import openpyxl

from app.services.parsing.base import Block, ParsedDocument, ParsingError


def parse_xlsx(content: bytes) -> ParsedDocument:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ParsingError(f"Failed to open XLSX: {exc}") from exc

    blocks: list[Block] = []
    has_data = False
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        blocks.append(Block(text=sheet.title, page_number=sheet_index, heading_level=1))
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cell.strip() for cell in cells):
                has_data = True
                blocks.append(Block(text=" | ".join(cells), page_number=sheet_index, heading_level=0))

    workbook.close()

    if not has_data:
        raise ParsingError("XLSX contains no extractable data")

    return ParsedDocument(title=None, blocks=blocks)
